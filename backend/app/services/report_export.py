"""
Section 57: Export Laporan — Excel & PDF for the 9 report types listed in the
spec (WhatsApp export lives separately in wa_generator.py, since it targets
chat text rather than a file).

Design: a single `build_report_rows()` dispatcher normalizes every report
into a common shape —

    {"title": str, "subtitle": str, "columns": [(key, label, fmt)], "rows": [dict]}

`fmt` is one of "text" | "number" | "currency" | "percent" | "date" and
drives both the Excel cell number_format and the PDF's Indonesian-style
text formatting (section 58), so a single builder feeds both renderers
without 18 separate endpoint implementations.

Excel cells keep raw numeric values (so the sheet is usable for further
analysis/pivoting in Excel) with a matching number_format applied. The PDF
renders section-58 Indonesian-formatted strings (Rp1,25 M / 84,5%), since a
PDF is a presentation artifact, not a raw-data one.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models import Pipeline
from app.services import funding_calc, id_format, monthly_report, pipeline_calc, target_calc

REPORT_TITLES = {
    "funding": "Posisi Funding per RMFT",
    "rmft_performance": "RMFT Performance Leaderboard",
    "pipeline": "Data Pipeline",
    "pipeline_conversion": "Tren Konversi Pipeline Harian",
    "customer_outflow": "Top Customer Outflow",
    "customer_inflow": "Top Customer Inflow",
    "target_achievement": "Target vs Achievement",
    "daily_performance": "Performance Harian RMFT",
    "monthly_performance": "Monthly Performance Summary",
}


def _today_month() -> str:
    return date.today().strftime("%Y-%m")


# ---------------------------------------------------------------------------
# Per-report builders — each returns the common {title, subtitle, columns, rows} shape.
# ---------------------------------------------------------------------------
def _funding(db: Session, params: dict) -> dict:
    pn = params.get("pn")
    data = funding_calc.compute_home_kpis(db)
    cards = data["rmft_cards"]
    if pn:
        cards = [c for c in cards if c["pn"] == pn]

    rows = [{
        "pn": c["pn"], "rmft_name": c["rmft_name"],
        "tabungan": c["current"]["tabungan"], "giro": c["current"]["giro"], "deposito": c["current"]["deposito"],
        "casa": c["current"]["casa"], "dpk": c["current"]["dpk"],
        "mtd_dpk": c["mtd"]["dpk"], "dtd_dpk": c["dtd"]["dpk"], "account_count": c["account_count"],
    } for c in cards]

    freshness = data["freshness"]
    as_of = id_format.format_date_id(freshness["as_of"]) if freshness["as_of"] else "-"
    return {
        "title": REPORT_TITLES["funding"],
        "subtitle": f"Data per {as_of} ({freshness['status']})",
        "columns": [
            ("pn", "PN", "text"), ("rmft_name", "RMFT", "text"),
            ("tabungan", "Tabungan", "currency"), ("giro", "Giro", "currency"), ("deposito", "Deposito", "currency"),
            ("casa", "CASA", "currency"), ("dpk", "Total DPK", "currency"),
            ("mtd_dpk", "MTD DPK", "currency"), ("dtd_dpk", "DTD DPK", "currency"),
            ("account_count", "Jumlah Rekening", "number"),
        ],
        "rows": rows,
    }


def _rmft_performance(db: Session, params: dict) -> dict:
    pn = params.get("pn")
    # date_from/date_to (from the frontend's Daily/MTD/YTD leaderboard scope selector)
    # take priority; otherwise fall back to a plain calendar month.
    if params.get("date_from") and params.get("date_to"):
        date_from, date_to = params["date_from"], params["date_to"]
        subtitle = f"Periode {id_format.format_date_id(date_from)} — {id_format.format_date_id(date_to)}"
    else:
        month = params.get("month") or _today_month()
        date_from, date_to = target_calc.month_range(month)
        subtitle = f"Periode {month}"
    rows = pipeline_calc.leaderboard(db, date_from=date_from, date_to=date_to)
    if pn:
        rows = [r for r in rows if r["pn"] == pn]
    return {
        "title": REPORT_TITLES["rmft_performance"],
        "subtitle": subtitle,
        "columns": [
            ("rank", "Rank", "number"), ("medal", "Medali", "text"),
            ("rmft_name", "RMFT", "text"), ("pn", "PN", "text"),
            ("pipeline_nominal", "Pipeline Nominal", "currency"), ("realisasi_nominal", "Realisasi Nominal", "currency"),
            ("outstanding", "Outstanding", "currency"),
            ("nominal_sr", "Success Rate Nominal", "percent"), ("activity_sr", "Success Rate Activity", "percent"),
            ("total_pipeline_count", "Jumlah Pipeline", "number"), ("realized_count", "Terealisasi", "number"),
            ("batal_count", "Batal", "number"),
        ],
        "rows": rows,
    }


def _daily_performance(db: Session, params: dict) -> dict:
    pn = params.get("pn")
    target_date: date = params.get("date") or date.today()
    rows = pipeline_calc.leaderboard(db, date_from=target_date, date_to=target_date)
    if pn:
        rows = [r for r in rows if r["pn"] == pn]
    return {
        "title": REPORT_TITLES["daily_performance"],
        "subtitle": f"Tanggal {id_format.format_date_id(target_date)}",
        "columns": [
            ("rank", "Rank", "number"), ("rmft_name", "RMFT", "text"), ("pn", "PN", "text"),
            ("pipeline_nominal", "Pipeline Nominal", "currency"), ("realisasi_nominal", "Realisasi Nominal", "currency"),
            ("nominal_sr", "Success Rate Nominal", "percent"), ("activity_sr", "Success Rate Activity", "percent"),
            ("total_pipeline_count", "Jumlah Pipeline", "number"), ("realized_count", "Terealisasi", "number"),
        ],
        "rows": rows,
    }


def _pipeline(db: Session, params: dict) -> dict:
    q = db.query(Pipeline)
    pn = params.get("pn")
    if pn:
        q = q.filter(Pipeline.pn == pn)
    date_from = params.get("date_from")
    date_to = params.get("date_to")
    if date_from:
        q = q.filter(Pipeline.pipeline_date >= date_from)
    if date_to:
        q = q.filter(Pipeline.pipeline_date <= date_to)
    status = params.get("status")
    if status:
        q = q.filter(Pipeline.status == status)

    rows = q.order_by(Pipeline.pipeline_date.desc()).limit(5000).all()
    row_dicts = [{
        "pipeline_date": r.pipeline_date, "pn": r.pn, "rmft": r.rmft, "customer": r.customer,
        "category": r.category, "product": r.product, "nominal": float(r.nominal or 0),
        "probability": r.probability, "target_date": r.target_date, "status": r.status,
    } for r in rows]

    return {
        "title": REPORT_TITLES["pipeline"],
        "subtitle": f"{len(row_dicts)} baris" + (f" — PN {pn}" if pn else ""),
        "columns": [
            ("pipeline_date", "Tanggal", "date"), ("pn", "PN", "text"), ("rmft", "RMFT", "text"),
            ("customer", "Customer", "text"), ("category", "Kategori", "text"), ("product", "Produk", "text"),
            ("nominal", "Nominal", "currency"), ("probability", "Probability (%)", "number"),
            ("target_date", "Target Tanggal", "date"), ("status", "Status", "text"),
        ],
        "rows": row_dicts,
    }


def _pipeline_conversion(db: Session, params: dict) -> dict:
    pn = params.get("pn")
    month = params.get("month") or _today_month()
    trend = monthly_report.daily_conversion_trend(db, month, pn=pn)
    return {
        "title": REPORT_TITLES["pipeline_conversion"],
        "subtitle": f"Periode {month}" + (f" — PN {pn}" if pn else ""),
        "columns": [
            ("date", "Tanggal", "date"), ("pipeline_nominal", "Pipeline Nominal", "currency"),
            ("realisasi_nominal", "Realisasi Nominal", "currency"),
            ("nominal_sr", "Success Rate Nominal", "percent"), ("activity_sr", "Success Rate Activity", "percent"),
        ],
        "rows": trend,
    }


def _customer_movers(db: Session, params: dict, direction: str) -> dict:
    pn = params.get("pn")
    mode = params.get("mode") or "MTD"
    limit = int(params.get("limit") or 20)
    rows = funding_calc.top_movers(db, mode=mode, direction=direction, pn=pn, limit=limit)
    key = "customer_outflow" if direction == "outflow" else "customer_inflow"
    return {
        "title": REPORT_TITLES[key],
        "subtitle": f"Mode {mode}" + (f" — PN {pn}" if pn else ""),
        "columns": [
            ("rank", "Rank", "number"), ("customer", "Customer", "text"), ("account_number", "No. Rekening", "text"),
            ("product", "Produk", "text"), ("rmft_name", "RMFT", "text"), ("pn", "PN", "text"),
            ("baseline_or_previous", "Saldo Referensi", "currency"), ("current_balance", "Saldo Saat Ini", "currency"),
            ("delta", "Delta", "currency"), ("delta_pct", "Delta (%)", "percent"), ("status", "Status", "text"),
        ],
        "rows": rows,
    }


def _target_achievement(db: Session, params: dict) -> dict:
    pn = params.get("pn")
    month = params.get("month") or _today_month()
    rows = target_calc.compute_achievement(db, month, pn=pn)
    return {
        "title": REPORT_TITLES["target_achievement"],
        "subtitle": f"Periode {month}",
        "columns": [
            ("rmft_name", "RMFT", "text"), ("pn", "PN", "text"), ("product", "Produk", "text"),
            ("target", "Target", "currency"), ("realisasi", "Realisasi", "currency"),
            ("achievement_pct", "Achievement (%)", "percent"), ("gap", "Gap", "currency"),
        ],
        "rows": rows,
    }


def _monthly_performance(db: Session, params: dict) -> dict:
    pn = params.get("pn")
    month = params.get("month") or _today_month()
    summary = monthly_report.monthly_summary(db, month, pn=pn)

    top_customer = summary.get("customer_terbesar") or {}
    top_product = summary.get("produk_terbesar") or {}
    rows = [
        {"metric": "Total Pipeline Nominal", "value": summary["pipeline_nominal"], "fmt": "currency"},
        {"metric": "Total Realisasi Nominal", "value": summary["realisasi_nominal"], "fmt": "currency"},
        {"metric": "Outstanding", "value": summary["outstanding"], "fmt": "currency"},
        {"metric": "Jumlah Pipeline", "value": summary["total_pipeline_count"], "fmt": "number"},
        {"metric": "Terealisasi", "value": summary["realized_count"], "fmt": "number"},
        {"metric": "Batal", "value": summary["batal_count"], "fmt": "number"},
        {"metric": "Success Rate Nominal", "value": summary["nominal_sr"], "fmt": "percent"},
        {"metric": "Success Rate Activity", "value": summary["activity_sr"], "fmt": "percent"},
        {"metric": "RMFT Terbaik", "value": summary.get("rmft_terbaik") or "-", "fmt": "text"},
        {"metric": "Customer Terbesar", "value": top_customer.get("customer") or "-", "fmt": "text"},
        {"metric": "Nominal Customer Terbesar", "value": top_customer.get("nominal") or 0, "fmt": "currency"},
        {"metric": "Produk Terbesar (Realisasi)", "value": top_product.get("product") or "-", "fmt": "text"},
        {"metric": "Realisasi Produk Terbesar", "value": top_product.get("realisasi") or 0, "fmt": "currency"},
    ]
    return {
        "title": REPORT_TITLES["monthly_performance"],
        "subtitle": f"Periode {month}" + (f" — PN {pn}" if pn else " — Seluruh Unit"),
        "columns": [("metric", "Metrik", "text"), ("value", "Nilai", "mixed")],
        "rows": rows,
    }


DISPATCH = {
    "funding": _funding,
    "rmft_performance": _rmft_performance,
    "daily_performance": _daily_performance,
    "pipeline": _pipeline,
    "pipeline_conversion": _pipeline_conversion,
    "customer_outflow": lambda db, p: _customer_movers(db, p, "outflow"),
    "customer_inflow": lambda db, p: _customer_movers(db, p, "inflow"),
    "target_achievement": _target_achievement,
    "monthly_performance": _monthly_performance,
}


def build_report_rows(db: Session, report: str, params: dict) -> dict:
    builder = DISPATCH.get(report)
    if not builder:
        raise ValueError(f"Tipe laporan tidak dikenal: {report}. Pilihan: {', '.join(DISPATCH.keys())}")
    return builder(db, params)


# ---------------------------------------------------------------------------
# Renderers
# ---------------------------------------------------------------------------
_EXCEL_NUMBER_FORMATS = {
    "currency": "#,##0",
    "number": "#,##0",
    "percent": '0.0"%"',
    "date": "dd/mm/yyyy",
    "text": "@",
}


def to_excel(result: dict) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = result["title"][:31] or "Report"

    columns = result["columns"]
    ws.append([result["title"]])
    ws.append([result["subtitle"]])
    ws.append([f"Diekspor: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    ws.append([])
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"].font = Font(italic=True, color="666666")
    ws["A3"].font = Font(italic=True, size=9, color="999999")

    header_row_idx = 5
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0B2545", end_color="0B2545", fill_type="solid")
    for col_idx, (key, label, fmt) in enumerate(columns, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_offset, row in enumerate(result["rows"], start=1):
        excel_row = header_row_idx + row_offset
        for col_idx, (key, label, fmt) in enumerate(columns, start=1):
            # "monthly_performance" rows carry a per-row fmt override (mixed metric list)
            cell_fmt = row.get("fmt", fmt) if fmt == "mixed" else fmt
            value = row.get(key)
            cell = ws.cell(row=excel_row, column=col_idx)
            if value is None:
                cell.value = "-"
            elif cell_fmt in ("currency", "number", "percent"):
                try:
                    cell.value = float(value)
                except (TypeError, ValueError):
                    cell.value = str(value)
                else:
                    cell.number_format = _EXCEL_NUMBER_FORMATS.get(cell_fmt, "General")
            elif cell_fmt == "date":
                cell.value = value
                cell.number_format = _EXCEL_NUMBER_FORMATS["date"]
            else:
                cell.value = str(value)

    for col_idx in range(1, len(columns) + 1):
        letter = get_column_letter(col_idx)
        max_len = max(
            [len(str(result["rows"][i].get(columns[col_idx - 1][0], "")) or "") for i in range(len(result["rows"]))]
            + [len(columns[col_idx - 1][1])]
        )
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _pdf_format(value, fmt: str) -> str:
    if value is None:
        return "-"
    if fmt == "currency":
        return id_format.format_short(value)
    if fmt == "percent":
        return id_format.format_percent(value)
    if fmt == "number":
        try:
            return f"{int(value):,}".replace(",", ".")
        except (TypeError, ValueError):
            return str(value)
    if fmt == "date":
        return id_format.format_date_id(value)
    return str(value)


def to_pdf(result: dict) -> io.BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    columns = result["columns"]
    page_size = landscape(A4) if len(columns) > 6 else A4
    doc = SimpleDocTemplate(buf, pagesize=page_size, topMargin=14 * mm, bottomMargin=14 * mm)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(result["title"], styles["Title"]),
        Paragraph(result["subtitle"], styles["Normal"]),
        Paragraph(f"Diekspor: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]),
        Spacer(1, 10),
    ]

    header = [label for _, label, _ in columns]
    table_data = [header]
    for row in result["rows"]:
        line = []
        for key, label, fmt in columns:
            cell_fmt = row.get("fmt", fmt) if fmt == "mixed" else fmt
            line.append(_pdf_format(row.get(key), cell_fmt))
        table_data.append(line)

    if len(table_data) == 1:
        elements.append(Paragraph("Tidak ada data untuk kriteria ini.", styles["Normal"]))
    else:
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b2545")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f5f9")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(table)

    doc.build(elements)
    buf.seek(0)
    return buf
