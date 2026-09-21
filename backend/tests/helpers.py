import openpyxl


def build_di319(path, periode="20/08/2026", date_printed="21/08/2026", extra_rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["BANK XYZ"])
    ws.append(["SAVINGS ACCOUNT MONTHLY TRIAL BALANCE"])
    ws.append(["Date Printed:", date_printed])
    ws.append([])
    ws.append([
        "periode", "account number", "ciff no", "short name", "prod code", "curr code",
        "balance", "Balance dalam IDR", "PN RM Dana/Mantri", "PN RM Referral", "PN Customer Service",
    ])
    ws.append([periode, "1234567890", "CIF001", "PT ABC", "SAV", "IDR",
               1_000_000, 1_000_000, "00382271 Ahmad Rafiq", "00274689 - Dia Silopa Putri", None])
    ws.append([periode, "1234567891", "CIF002", "PT DEF", "SAV", "IDR",
               2_000_000, 2_000_000, "00380727 - Adist Ayudistira Sembiring", None, None])
    for extra in (extra_rows or []):
        ws.append(extra)
    ws.append(["", "TOTAL", "", "", "", "", 3000000, 3000000, "", "", ""])
    wb.save(path)


def build_di321(path, periode="20/08/2026"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["CURRENT ACCOUNT MONTHLY TRIAL BALANCE"])
    ws.append(["Date Printed", "21/08/2026"])
    ws.append([])
    ws.append([
        "PERIODE", "CIFNO", "ACCOUNT NUMBER", "PRODUCT CODE", "SHORT NAME", "STATUS",
        "BALANCE", "AVAIL BALANCE", "AVRG BALANCE", "CURR", "PN RM DANA",
    ])
    ws.append([periode, "CIF010", "9988776655", "GIR", "PT SUKSES", "ACTIVE",
               15_000_000, 14_500_000, 14_800_000, "IDR", "00274689 Dia Silopa"])
    wb.save(path)


def build_ci324(path, periode="20/08/2026"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["FDs MONTHLY TRIAL BALANCE"])
    ws.append(["Date Printed :", "21/08/2026"])
    ws.append([])
    ws.append([
        "PERIODE", "ACCTNO", "FDR SRL NO", "SHORT NAME", "TYPE", "PRINCIPAL AMOUNT",
        "ISSUE DT", "MAT DT", "INT RATE", "INT TENOR DISP", "RENEW", "CBAL Base", "PN RM Dana",
    ])
    ws.append([periode, "5566778899", "FDR001", "PT MAKMUR", "FD", 100_000_000,
               "20/02/2026", "20/02/2027", 5.5, "12M", "AUTO", None, "00380727 Adist Ayudistira"])
    wb.save(path)
